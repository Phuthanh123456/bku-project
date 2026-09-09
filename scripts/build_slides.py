"""Sinh bộ slide bảo vệ 15 trang từ số liệu/ảnh thật trong results/."""

from __future__ import annotations

import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

GOC = Path(__file__).resolve().parents[1]
OUTPUT = GOC / "results/ENVI_NMT_bao_ve.pptx"
BLUE = RGBColor(36, 84, 160)
DARK = RGBColor(28, 37, 54)
MUTED = RGBColor(92, 103, 125)
for _stream in (sys.stdout, sys.stderr):
    if hasattr(_stream, "reconfigure"):
        _stream.reconfigure(encoding="utf-8", errors="replace")


def _style_text_frame(frame, size=22, color=DARK):
    for paragraph in frame.paragraphs:
        paragraph.font.name = "Aptos"
        paragraph.font.size = Pt(size)
        paragraph.font.color.rgb = color
        paragraph.space_after = Pt(9)


def _title(slide, title: str, number: int):
    box = slide.shapes.add_textbox(Inches(0.6), Inches(0.25), Inches(11.9), Inches(0.65))
    p = box.text_frame.paragraphs[0]
    p.text = title
    p.font.name = "Aptos Display"
    p.font.bold = True
    p.font.size = Pt(30)
    p.font.color.rgb = BLUE
    footer = slide.shapes.add_textbox(Inches(12.25), Inches(7.05), Inches(0.45), Inches(0.25))
    fp = footer.text_frame.paragraphs[0]
    fp.text = str(number)
    fp.font.size = Pt(10)
    fp.font.color.rgb = MUTED
    fp.alignment = PP_ALIGN.RIGHT


def _bullet_slide(prs, title, bullets, number):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _title(slide, title, number)
    box = slide.shapes.add_textbox(Inches(0.85), Inches(1.2), Inches(11.4), Inches(5.5))
    frame = box.text_frame
    frame.clear()
    for index, item in enumerate(bullets):
        p = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        p.text = item
        p.level = 0
    _style_text_frame(frame, 23)
    return slide


def _image_slide(prs, title, image_name, caption, number):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _title(slide, title, number)
    image_path = GOC / "results" / image_name
    slide.shapes.add_picture(str(image_path), Inches(0.8), Inches(1.0), width=Inches(11.7), height=Inches(5.65))
    cap = slide.shapes.add_textbox(Inches(0.9), Inches(6.75), Inches(11.4), Inches(0.35))
    cap.text_frame.text = caption
    _style_text_frame(cap.text_frame, 13, MUTED)
    return slide


def _table_slide(prs, title, headers, rows, number, note=""):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _title(slide, title, number)
    table = slide.shapes.add_table(
        len(rows) + 1, len(headers), Inches(0.65), Inches(1.25), Inches(12.0), Inches(4.8)
    ).table
    for col, header in enumerate(headers):
        table.cell(0, col).text = header
    for r, row in enumerate(rows, start=1):
        for c, value in enumerate(row):
            table.cell(r, c).text = str(value)
    for r in range(len(rows) + 1):
        for c in range(len(headers)):
            cell = table.cell(r, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = BLUE if r == 0 else RGBColor(242, 246, 252)
            _style_text_frame(cell.text_frame, 16, RGBColor(255, 255, 255) if r == 0 else DARK)
    if note:
        box = slide.shapes.add_textbox(Inches(0.8), Inches(6.3), Inches(11.7), Inches(0.55))
        box.text_frame.text = note
        _style_text_frame(box.text_frame, 15, MUTED)
    return slide


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    title = slide.shapes.add_textbox(Inches(0.8), Inches(1.55), Inches(11.8), Inches(1.3))
    title.text_frame.text = "ENVI-NMT: Transformer Anh→Việt tự xây dựng"
    _style_text_frame(title.text_frame, 38, BLUE)
    subtitle = slide.shapes.add_textbox(Inches(0.85), Inches(3.05), Inches(11.5), Inches(1.5))
    subtitle.text_frame.text = "RoPE · Pre-Norm · RMSNorm · SwiGLU\nIWSLT 2015 · Cập nhật 08/09/2026"
    _style_text_frame(subtitle.text_frame, 24, MUTED)

    _bullet_slide(prs, "1. Bài toán", [
        "Dịch máy Anh→Việt trong bối cảnh dữ liệu ít (IWSLT 2015).",
        "Tự cài Transformer bằng PyTorch; không dùng model pretrained/nn.Transformer.",
        "Yêu cầu: kết quả tái lập, checkpoint chịu lỗi và demo sử dụng được.",
    ], 2)
    _bullet_slide(prs, "2. Câu hỏi nghiên cứu & target", [
        "Cải tiến RoPE + Pre-Norm + RMSNorm + SwiGLU có hơn vanilla 2017 cùng budget?",
        "Cổng chất lượng: tst2013 BLEU ≥ 22; giả thuyết đóng góp +1,0 BLEU.",
        "Mọi ablation chạy đúng 6.000 bước, hai seed 42 và 1337.",
    ], 3)
    _table_slide(prs, "3. Dữ liệu sau làm sạch", ["Split", "Trước", "Sau", "Vai trò"], [
        ["train", "133.317", "131.339", "huấn luyện"],
        ["tst2012", "1.553", "1.553", "dev"],
        ["tst2013", "1.268", "1.268", "test"],
    ], 4, "Đã loại 21 cặp rò rỉ train↔dev/test; test không dùng chọn model.")
    _bullet_slide(prs, "4. Tokenizer", [
        "BPE dùng chung Anh–Việt, vocab 32.000, embedding chia sẻ.",
        "Fertility: EN 1,0664 · VI 1,0137; UNK train/dev = 0,0000%.",
        "Corpus đã tách token → điểm chính dùng SacreBLEU tok:none.",
    ], 5)
    _table_slide(prs, "5. Hyperparameter cuối", ["Thành phần", "Giá trị"], [
        ["Encoder / Decoder", "6 / 6 lớp"], ["Attention", "d_model 512; 8 head × 64"],
        ["FFN", "SwiGLU d_ff 688"], ["Dropout", "0,3"],
        ["Optimizer", "AdamW lr 7e-4; clip 1,0"], ["Batch", "4.096 token × tích lũy 4"],
    ], 6)
    _bullet_slide(prs, "6. Bốn thay đổi kiến trúc", [
        "RoPE: áp query/key của self-attention; không áp cross-attention.",
        "Pre-Norm + final RMSNorm: ổn định đường gradient.",
        "RMSNorm: chuẩn hóa độ lớn, tính mean-square ở float32.",
        "SwiGLU d_ff=688 khớp tham số ReLU d_ff=1024 dưới 1%.",
    ], 7)
    _bullet_slide(prs, "7. Kiểm chứng tính đúng", [
        "Attention/RMSNorm đối chiếu PyTorch sai lệch < 1e-5.",
        "Beam=1 trùng Greedy; cache trùng decode đầy đủ cho RoPE và sin-cos.",
        "Mô hình 47.955.968 tham số; vanilla khớp cỡ chênh 0,175%.",
        "Cổng overfit 50 câu nằm trong test slow của repo.",
    ], 8)
    _image_slide(prs, "8. Checkpoint chịu lỗi", "thi_nghiem_phuc_hoi.png", "60/60 bước có loss trùng; chênh lớn nhất 0,000000%.", 9)
    _image_slide(prs, "9. Huấn luyện", "training_curves.png", "Checkpoint được chấm: bước 6.000, epoch 34, loss dev 2,2364.", 10)
    _table_slide(prs, "10. Kết quả full benchmark", ["Split", "Search", "BLEU", "chrF++"], [
        ["tst2012", "Greedy", "26,24", "45,57"],
        ["tst2013", "Greedy", "29,74", "48,73"],
        ["tst2013", "Beam-4", "30,69", "49,60"],
    ], 11, "Signature BLEU: tok:none · Beam tăng +0,95 BLEU.")
    _image_slide(prs, "11. Greedy vs Beam", "search_comparison.png", "Beam cải thiện chất lượng corpus nhưng chậm hơn; không sửa mọi lỗi từng câu.", 12)
    _image_slide(prs, "12. Hiệu quả KV-cache", "latency_comparison.png", "200 câu CPU: throughput 0,77→2,41 câu/s; BLEU giữ nguyên.", 13)
    _image_slide(prs, "13. Phân tích lỗi 30 câu", "phan_tich_loi.png", "Nổi bật: sai nghĩa từ đa nghĩa và dịch sát cấu trúc tiếng Anh.", 14)
    _bullet_slide(prs, "14. Kết luận & việc còn thiếu", [
        "Đã đạt chất lượng: Greedy 29,74; Beam-4 30,69 BLEU.",
        "Demo Docker đang chạy local; image 429 MiB, health HTTP 200.",
        "Chưa có 15 lượt GPU vanilla/A1–A6 còn lại → chưa kết luận đóng góp nhân quả.",
        "Tiếp theo: chạy manifest trên T4, tổng hợp mean/std, tạo URL Space bằng tài khoản nhóm.",
    ], 15)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUTPUT)
    print(f"Đã lưu {OUTPUT}")


if __name__ == "__main__":
    main()
